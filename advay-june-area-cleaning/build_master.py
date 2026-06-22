# -*- coding: utf-8 -*-
"""
Build a 'Master Data' sheet (Project + Tower + Flat-line -> carpet areas) and
backfill Index II blanks that have a consistent internal match.

Lines with NO extracted data anywhere (e.g. Skyline Icon, Lodha Altus A&B wings)
are listed with Source='NEEDS EXTERNAL DATA' so the external competition master
can be merged in later.
"""
import re
import numpy as np
import pandas as pd
import openpyxl

WB = "/home/user/rajw56/advay-june-area-cleaning/Advay_June_All_projects.xlsx"
SQM2SQFT = 10.764
SAFE_SPREAD = 10.0   # sqft: only auto-backfill a blank if its line is this consistent


def flat_line(fn):
    s = re.sub(r'\D', '', str(fn))
    return s[-2:] if len(s) >= 2 else (s or '?')


def main():
    df = pd.read_excel(WB, sheet_name='Index II')
    act = pd.to_numeric(df['Actual Areas (Sqft)'], errors='coerce')
    flat = pd.to_numeric(df['Flat Area (Sqm)'], errors='coerce')
    balc = pd.to_numeric(df['Balcony area (Sqm)'], errors='coerce')
    add = pd.to_numeric(df['Additional Area (Sqm)'], errors='coerce')

    df['proj'] = df['Project Name- As per Our Database'].astype(str)
    df['tower'] = df['Tower Name'].fillna('-').astype(str)
    df['line'] = df['Flat No'].apply(flat_line)

    # ---------- build master ----------
    master = []
    for (p, t, ln), g in df.groupby(['proj', 'tower', 'line']):
        gi = g.index
        a = act.loc[gi].dropna()
        sample = ', '.join(sorted(g['Flat No'].dropna().astype(str).unique())[:6])
        if len(a):
            spread = round(a.max() - a.min(), 2)
            master.append(dict(
                Project=p, Tower=(None if t == '-' else t), **{'Flat Line': ln},
                **{'Flat Area (Sqm)': round(flat.loc[gi].dropna().median(), 2),
                   'Balcony area (Sqm)': (round(balc.loc[gi].dropna().median(), 2) if balc.loc[gi].notna().any() else None),
                   'Additional Area (Sqm)': (round(add.loc[gi].dropna().median(), 2) if add.loc[gi].notna().any() else None),
                   'Actual Areas (Sqft)': round(a.median(), 2),
                   'Source': 'Extracted from descriptions',
                   'Transactions': int(len(a)),
                   'Area Spread (Sqft)': spread,
                   'Sample Flat Nos': sample}))
        else:
            master.append(dict(
                Project=p, Tower=(None if t == '-' else t), **{'Flat Line': ln},
                **{'Flat Area (Sqm)': None, 'Balcony area (Sqm)': None,
                   'Additional Area (Sqm)': None, 'Actual Areas (Sqft)': None,
                   'Source': 'NEEDS EXTERNAL DATA',
                   'Transactions': 0, 'Area Spread (Sqft)': None,
                   'Sample Flat Nos': sample}))
    mdf = pd.DataFrame(master).sort_values(['Project', 'Tower', 'Flat Line'], na_position='first')

    # lookup for safe backfill: (proj,tower,line) -> medians, where extracted & consistent
    look = {}
    for r in master:
        if r['Source'].startswith('Extracted') and (r['Area Spread (Sqft)'] or 0) <= SAFE_SPREAD:
            look[(r['Project'], r['Tower'] or '-', r['Flat Line'])] = r

    # ---------- write with openpyxl (preserve Index II, add Master Data) ----------
    wb = openpyxl.load_workbook(WB)
    ws = wb['Index II']
    H = {c.value: c.column for c in ws[1]}

    filled_n = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=H['Actual Areas (Sqft)']).value not in (None, ''):
            continue
        p = str(ws.cell(row=row, column=H['Project Name- As per Our Database']).value)
        tv = ws.cell(row=row, column=H['Tower Name']).value
        t = '-' if tv in (None, '') else str(tv)
        ln = flat_line(ws.cell(row=row, column=H['Flat No']).value)
        ref = look.get((p, t, ln))
        if ref:
            ws.cell(row=row, column=H['Flat Area (Sqm)']).value = ref['Flat Area (Sqm)']
            ws.cell(row=row, column=H['Balcony area (Sqm)']).value = ref['Balcony area (Sqm)']
            ws.cell(row=row, column=H['Additional Area (Sqm)']).value = ref['Additional Area (Sqm)']
            ws.cell(row=row, column=H['Actual Areas (Sqft)']).value = ref['Actual Areas (Sqft)']
            ws.cell(row=row, column=H['Area and Price Remarks']).value = \
                'filled from Master Data (%s | Tower %s | line %s)' % (p, t, ln)
            filled_n += 1

    if 'Master Data' in wb.sheetnames:
        del wb['Master Data']
    ms = wb.create_sheet('Master Data')
    cols = ['Project', 'Tower', 'Flat Line', 'Flat Area (Sqm)', 'Balcony area (Sqm)',
            'Additional Area (Sqm)', 'Actual Areas (Sqft)', 'Source', 'Transactions',
            'Area Spread (Sqft)', 'Sample Flat Nos']
    ms.append(cols)
    for _, r in mdf.iterrows():
        ms.append([r[c] for c in cols])

    wb.save(WB)

    # ---------- report ----------
    tot = len(mdf)
    ext = (mdf['Source'] == 'Extracted from descriptions').sum()
    need = (mdf['Source'] == 'NEEDS EXTERNAL DATA').sum()
    print('Master Data rows:', tot, '| Extracted:', ext, '| NEEDS EXTERNAL:', need)
    print('Index II blanks backfilled from master:', filled_n)
    print('\nLines NEEDING external data (by project):')
    print(mdf[mdf.Source == 'NEEDS EXTERNAL DATA'].groupby('Project')['Flat Line']
          .agg(['count']).to_string())
    rem_blank = pd.read_excel(WB, sheet_name='Index II')['Actual Areas (Sqft)'].isna().sum()
    print('\nIndex II rows still blank:', rem_blank)


if __name__ == '__main__':
    main()
